import logging
import re
from pathlib import Path

import docx
from pypdf import PdfReader

from app.config import settings
from app.errors import AppError
from app.models import FileType

logger = logging.getLogger("app.files")

TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "cp1251")


class ExtractError(AppError):
    """Из файла не удалось достать текст."""
    def __init__(self, message: str, code: str = "extract_failed") -> None:
        super().__init__(message, status_code=422, code=code)


def clean_text(raw: str) -> str:
    """Привести текст к аккуратному виду"""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: Path) -> str:
    """Достать текст из PDF постранично."""
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # битый или зашифрованный файл
        raise ExtractError(f"Не удалось открыть PDF: {exc}", "bad_pdf") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractError("PDF защищён паролем", "encrypted_pdf") from exc

    pages: list[str] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            logger.warning("Страница %s не прочиталась: %s", number, path.name)
            pages.append("")
    return "\n\n".join(pages)


def extract_docx(path: Path) -> str:
    """Достать текст из DOCX: абзацы и содержимое таблиц."""
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractError(
            "Не удалось открыть DOCX. Возможно, это файл другого формата "
            "с расширением .docx",
            "bad_docx",
        ) from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_txt(path: Path) -> str:
    """Прочитать TXT, подбирая кодировку."""
    data = path.read_bytes()
    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractError(
        "Не удалось определить кодировку файла. "
        "Пересохраните его в UTF-8",
        "bad_encoding",
    )


EXTRACTORS = {
    FileType.PDF: extract_pdf,
    FileType.DOCX: extract_docx,
    FileType.TXT: extract_txt,
}


def extract_text(path: Path, file_type: FileType) -> str:
    text = clean_text(EXTRACTORS[file_type](path))

    if len(text) < settings.min_text_chars:
        raise ExtractError(
            "В файле не нашлось текста. Скорее всего, это скан или фотографии "
            "страниц: внутри такого файла только картинки. "
            "Загрузите документ с текстовым слоем",
            "no_text_layer",
        )
    logger.info("Извлечено %s символов из %s", len(text), path.name)
    return text